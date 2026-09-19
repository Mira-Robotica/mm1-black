#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Termo de Empréstimo — MM1-BLACK (Trena 2D)
Validação por espeleólogos · Timbre MIRA (estilo pré-propostas ERO)
"""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DIR = Path(__file__).resolve().parent
OUT = DIR / "Termo_Emprestimo_MM1-BLACK.docx"
FIG = DIR / "assets" / "figures"
LOGO = FIG / "mira-horizontal.png"
if not LOGO.is_file():
    LOGO = FIG / "mira-logo.png"
LOGO_MARK = FIG / "mira-principal.png"

DARK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x5A, 0x5A, 0x5A)
MIRA_RED = RGBColor(0xC8, 0x20, 0x2F)
CHARCOAL = "404040"
LINE = "C8C8C8"
BOX_FILL = "F4F4F4"
HEAD_FILL = "ECECEC"

DOC_CODE = "TE-MIRA-MM1-001"
REV = "A"
TODAY = date.today().strftime("%d/%m/%Y")


def set_run_font(run, name="Arial", size=10, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell, color=LINE, sz="4"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for m, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        n = OxmlElement(f"w:{m}")
        n.set(qn("w:w"), str(v))
        n.set(qn("w:type"), "dxa")
        mar.append(n)
    tcPr.append(mar)


def add_bottom_border(paragraph, color=CHARCOAL, sz="12"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_page_number(paragraph):
    def fld(instr):
        run1 = paragraph.add_run()
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), "begin")
        run1._r.append(fc)
        run2 = paragraph.add_run()
        t = OxmlElement("w:instrText")
        t.set(qn("xml:space"), "preserve")
        t.text = instr
        run2._r.append(t)
        run3 = paragraph.add_run()
        fc2 = OxmlElement("w:fldChar")
        fc2.set(qn("w:fldCharType"), "end")
        run3._r.append(fc2)
        for run in (run1, run2, run3):
            set_run_font(run, size=8, color=GRAY)

    fld(" PAGE ")
    r = paragraph.add_run(" de ")
    set_run_font(r, size=8, color=GRAY)
    fld(" NUMPAGES ")


def para(doc, text, size=10, bold=False, italic=False, align="justify",
         space_after=8, space_before=0, color=DARK):
    p = doc.add_paragraph()
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.1
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return p


def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_run_font(run, size=12.5, bold=True, color=DARK)
    add_bottom_border(p, CHARCOAL, "10")
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=10.5, bold=True, color=DARK)
    return p


def info_box(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    shade(cell, BOX_FILL)
    set_cell_border(cell, CHARCOAL, "8")
    set_cell_margins(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_run_font(r, size=10, bold=True, color=DARK)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    set_run_font(r2, size=9.5, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def field_row(doc, label, blank="________________________________"):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(label + " ")
    set_run_font(r, size=10, bold=True, color=DARK)
    r2 = p.add_run(blank)
    set_run_font(r2, size=10, color=GRAY)


def setup_header_footer(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.9)
    section.different_first_page_header_footer = True

    # First page footer only (cover-like)
    fp = section.first_page_footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(f"{DOC_CODE} · Rev. {REV} · Confidencial")
    set_run_font(r, size=8, color=GRAY)

    # Header pages 2+
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = header.add_run("MIRA — Mapeamento e Robótica Ltda.")
    set_run_font(r, size=8, bold=True, color=DARK)
    r = header.add_run(f"                    {DOC_CODE} · Termo de Empréstimo MM1-BLACK · Confidencial")
    set_run_font(r, size=8, color=GRAY)
    add_bottom_border(header, CHARCOAL, "6")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = footer.add_run(f"Rev. {REV} · {TODAY}          ")
    set_run_font(r, size=8, color=GRAY)
    add_page_number(footer)
    r = footer.add_run("          Prazo típico: 3 a 6 meses")
    set_run_font(r, size=8, color=GRAY)
    pPr = footer._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "6")
    top.set(qn("w:space"), "1")
    top.set(qn("w:color"), CHARCOAL)
    pBdr.append(top)
    pPr.append(pBdr)


def cover_timbre(doc):
    """Logo + title block like ERO pré-proposta."""
    if LOGO.is_file():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run()
        run.add_picture(str(LOGO), width=Cm(7.0))
    elif LOGO_MARK.is_file():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(LOGO_MARK), width=Cm(3.5))

    para(doc, "MIRA — MAPEAMENTO E ROBÓTICA LTDA.", size=11, bold=True,
         align="center", space_after=2, color=MIRA_RED)
    para(doc, "CNPJ 67.815.751/0001-37 · Belo Horizonte/MG", size=9,
         align="center", space_after=12, color=GRAY)

    para(doc, "TERMO DE EMPRÉSTIMO DE EQUIPAMENTO", size=16, bold=True,
         align="center", space_after=4)
    para(doc, "MM1-BLACK — Trena 2D", size=13, bold=True,
         align="center", space_after=4, color=DARK)
    para(doc, "Destinado à validação técnica por espeleólogos", size=10, italic=True,
         align="center", space_after=6, color=GRAY)
    para(doc, f"Documento {DOC_CODE}  ·  Revisão {REV}  ·  {TODAY}", size=9,
         align="center", space_after=14, color=GRAY)


def build():
    doc = Document()
    setup_header_footer(doc)
    cover_timbre(doc)

    info_box(
        doc,
        "Objeto deste termo",
        "Empréstimo gratuito e temporário da trena MM1-BLACK (Trena 2D), unidade vedada "
        "com cartão de memória já preparado, para uso exclusivo em atividades de "
        "validação técnica, testes de campo e avaliação por espeleólogos e grupos "
        "parceiros, nas condições abaixo.",
    )

    # 1. Partes
    h1(doc, "1. Partes")
    para(
        doc,
        "1.1. COMODANTE: MIRA — Mapeamento e Robótica Ltda., pessoa jurídica de direito "
        "privado, inscrita no CNPJ sob o nº 67.815.751/0001-37, com sede em Belo Horizonte/MG, "
        "neste ato representada por seu representante legal, doravante denominada simplesmente MIRA.",
    )
    para(doc, "1.2. COMODATÁRIO (preencher):", bold=True, space_after=4)
    field_row(doc, "Nome / Razão social:")
    field_row(doc, "CPF / CNPJ:")
    field_row(doc, "Endereço:")
    field_row(doc, "E-mail / Telefone:")
    field_row(doc, "Grupo / Clube espeleológico (se houver):")
    para(
        doc,
        "1.3. As partes acima são doravante designadas, em conjunto, como Partes e, "
        "individualmente, como Parte.",
        space_before=6,
    )

    # 2. Objeto
    h1(doc, "2. Objeto do empréstimo")
    para(
        doc,
        "2.1. A MIRA cede em comodato (empréstimo gratuito), nos termos dos arts. 579 e "
        "seguintes do Código Civil brasileiro, o seguinte bem móvel:",
    )
    info_box(
        doc,
        "Equipamento",
        "Descrição: MM1-BLACK — Trena 2D (unidade vedada, com cartão microSD incluso e "
        "preparado de fábrica, firmware instalado).\n"
        "Número de série / patrimônio: _______________________________\n"
        "Acessórios entregues: _______________________________________________\n"
        "Estado aparente na entrega: ( ) íntegro  ( ) observações: _______________",
    )
    para(
        doc,
        "2.2. O equipamento permanece de propriedade exclusiva da MIRA durante toda a "
        "vigência deste Termo. O presente instrumento não transfere domínio, posse "
        "ad usucapionem, licença de exploração comercial nem qualquer direito real sobre "
        "o bem ou sobre a propriedade intelectual a ele associada.",
    )
    para(
        doc,
        "2.3. A finalidade do empréstimo é exclusiva e estritamente limitada à validação "
        "técnica, testes de campo, levantamento experimental e avaliação operacional por "
        "espeleólogos, não constituindo venda, locação onerosa, parceria comercial ou "
        "autorização para uso profissional remunerado perante terceiros, salvo autorização "
        "prévia e escrita da MIRA.",
    )

    # 3. Prazo
    h1(doc, "3. Prazo")
    para(
        doc,
        "3.1. O prazo típico de empréstimo é de 3 (três) a 6 (seis) meses, contados da "
        "data de entrega do equipamento, conforme preenchimento abaixo:",
    )
    field_row(doc, "Data de início (entrega):", "____/____/________")
    field_row(doc, "Data de término (devolução):", "____/____/________")
    field_row(doc, "Prazo total acordado:", "______ meses (entre 3 e 6)")
    para(
        doc,
        "3.2. A prorrogação somente será válida se acordada por escrito (inclusive "
        "correio eletrônico) pela MIRA, antes do vencimento. O silêncio da MIRA não "
        "implica renovação automática.",
        space_before=6,
    )
    para(
        doc,
        "3.3. Ao término do prazo, ou quando solicitado pela MIRA com aviso prévio de "
        "15 (quinze) dias, o Comodatário devolverá o equipamento no estado em que o "
        "recebeu, ressalvado o desgaste natural de uso adequado, no local e forma "
        "indicados pela MIRA.",
    )

    # 4. Obrigações e restrições
    h1(doc, "4. Obrigações e restrições de uso")
    para(doc, "4.1. O Comodatário obriga-se a:", bold=True, space_after=4)
    bullets = [
        "utilizar o equipamento exclusivamente para a finalidade prevista na cláusula 2.3;",
        "guardar o bem com o mesmo cuidado que empregaria se fosse próprio, adotando "
        "práticas seguras em campo (inclusive quanto a laser e ambientes confinados);",
        "não ceder, emprestar, sublocar, penhorar, dar em garantia ou de qualquer forma "
        "transferir o equipamento a terceiros sem autorização escrita da MIRA;",
        "comunicar imediatamente à MIRA qualquer dano, extravio, furto, avaria ou "
        "mau funcionamento;",
        "cumprir as contrapartidas da cláusula 5.",
    ]
    for b in bullets:
        para(doc, "• " + b, space_after=3)

    para(doc, "4.2. É expressamente vedado ao Comodatário:", bold=True,
         space_before=8, space_after=4)
    vedado = [
        "abrir, deslacrar, desmontar, furar, soldar, modificar a vedação, o invólucro, "
        "o hardware, o firmware ou quaisquer componentes internos do dispositivo;",
        "remover, substituir ou adulterar o cartão de memória, etiquetas, lacres, "
        "números de série ou identificação patrimonial;",
        "realizar engenharia reversa, cópia, extração ou divulgação de código-fonte, "
        "firmware, esquemas, desenhos ou know-how associados ao equipamento;",
        "vender, doar, ceder onerosamente, alugar, licenciar, comercializar ou explorar "
        "economicamente o dispositivo, seus acessórios ou qualquer parte deles;",
        "vender, doar, ceder, publicar comercialmente, sublicenciar ou de qualquer forma "
        "comercializar dados brutos ou derivados capturados pelo dispositivo, logs, "
        "arquivos de levantamento, nuvens de pontos, trajetórias, métricas, imagens de "
        "tela ou quaisquer informações geradas ou relacionadas ao uso do MM1-BLACK, "
        "salvo o compartilhamento com a MIRA previsto na cláusula 5 e usos acadêmicos "
        "não comerciais previamente autorizados por escrito;",
        "utilizar o equipamento de modo a gerar risco injustificado a pessoas, ao "
        "patrimônio espeleológico ou a terceiros, ou em desacordo com normas de "
        "segurança aplicáveis a espaços confinados.",
    ]
    for i, b in enumerate(vedado, 1):
        para(doc, f"({chr(96+i)}) {b}", space_after=3)

    para(
        doc,
        "4.3. A violação de qualquer vedação desta cláusula constitui infração grave, "
        "autorizando a MIRA a exigir a devolução imediata do equipamento, sem prejuízo "
        "de perdas e danos, medidas judiciais e responsabilização civil e, se cabível, "
        "criminal.",
        space_before=6,
    )

    # 5. Contrapartidas
    h1(doc, "5. Contrapartidas do empréstimo")
    para(
        doc,
        "5.1. Em contraprestação ao empréstimo gratuito, o Comodatário compromete-se a:",
    )
    para(
        doc,
        "(a) Divulgação: mencionar e divulgar a trena MM1-BLACK / Trena 2D e a marca "
        "MIRA em canais razoáveis de comunicação do Comodatário (por exemplo: redes "
        "sociais, grupos de espeleologia, relatórios de campo, apresentações técnicas "
        "ou notas públicas), de forma leal e verdadeira, sem associar a MIRA a "
        "endossos não autorizados. A MIRA poderá fornecer texto e imagens de apoio.",
    )
    para(
        doc,
        "(b) Compartilhamento de dados: fornecer à MIRA cópia dos dados capturados "
        "com o equipamento (arquivos de levantamento, registros relevantes e "
        "informações de contexto de validação), exclusivamente para fins de "
        "comparação, análise técnica, melhoria do produto e validação científica/"
        "operacional, sem transferência de titularidade dos dados brutos de campo "
        "do Comodatário para exploração comercial por terceiros.",
    )
    para(
        doc,
        "5.2. A MIRA poderá utilizar os dados recebidos na forma da alínea (b) de "
        "modo agregado ou anonimizado em materiais técnicos, desde que não divulgue "
        "informações que identifiquem locais sensíveis ou dados pessoais sem base "
        "legal ou consentimento adequado, quando aplicável.",
    )
    para(
        doc,
        "5.3. O não cumprimento das contrapartidas, após notificação e prazo razoável "
        "para sanar, autoriza a MIRA a dar por rescindido o empréstimo e a exigir a "
        "devolução do equipamento.",
    )

    # 6. Propriedade intelectual
    h1(doc, "6. Propriedade intelectual e confidencialidade")
    para(
        doc,
        "6.1. Permanecem com a MIRA todos os direitos de propriedade intelectual "
        "sobre o equipamento, firmware, software, design, marcas, manuais e "
        "documentação. Nada neste Termo concede licença ampla além do uso "
        "temporário do bem físico para a finalidade da cláusula 2.3.",
    )
    para(
        doc,
        "6.2. O Comodatário manterá em sigilo informações técnicas não públicas "
        "obtidas em razão do empréstimo, pelo prazo de 3 (três) anos após a "
        "devolução, exceto se já forem de domínio público sem culpa sua ou se a "
        "divulgação for exigida por lei ou ordem judicial.",
    )

    # 7. Responsabilidade
    h1(doc, "7. Responsabilidade, risco e seguro")
    para(
        doc,
        "7.1. O Comodatário assume a guarda e o risco do equipamento desde a "
        "entrega até a devolução efetiva à MIRA, respondendo por perda, furto, "
        "extravio, dano ou deterioração decorrente de uso inadequado, negligência "
        "ou violação deste Termo, podendo a MIRA exigir reparação, reposição ou "
        "indenização pelo valor de reposição do bem.",
    )
    para(
        doc,
        "7.2. A MIRA entrega o equipamento em comodato para validação; não se "
        "obriga a resultado específico de levantamento, continuidade de suporte "
        "além do razoavelmente acordado, nem se responsabiliza por decisões "
        "operacionais, de segurança ou científicas tomadas pelo Comodatário com "
        "base nos dados obtidos.",
    )
    para(
        doc,
        "7.3. Em nenhuma hipótese a MIRA responderá por lucros cessantes, danos "
        "indiretos ou danos emergentes decorrentes do uso do equipamento, salvo "
        "dolo ou culpa grave comprovados.",
    )
    para(
        doc,
        "7.4. O uso em cavernas e espaços confinados é de inteira responsabilidade "
        "do Comodatário, que declara possuir capacitação e seguir protocolos de "
        "segurança adequados. O laser do equipamento não deve ser apontado para "
        "olhos de pessoas ou animais.",
    )

    # 8. Devolução
    h1(doc, "8. Devolução e rescisão")
    para(
        doc,
        "8.1. Na devolução, as Partes poderão lavrar termo de recebimento "
        "registrando o estado do bem. A ausência de ressalva no recebimento não "
        "afasta a responsabilidade por vícios ocultos decorrentes de mau uso.",
    )
    para(
        doc,
        "8.2. Este Termo poderá ser rescindido de pleno direito pela MIRA em caso "
        "de inadimplemento, uso fora da finalidade, violação das vedações, ou por "
        "conveniência, mediante aviso prévio de 15 (quinze) dias, sem ônus para a "
        "MIRA além da logística razoável de recolhimento previamente acordada.",
    )

    # 9. LGPD
    h1(doc, "9. Dados pessoais")
    para(
        doc,
        "9.1. Dados pessoais eventualmente tratados em razão deste Termo serão "
        "tratados conforme a Lei nº 13.709/2018 (LGPD), para execução deste "
        "instrumento e legítimos interesses ligados à validação do produto, "
        "podendo o titular exercer seus direitos pelos canais da MIRA "
        "(contato@mirarobotica.com).",
    )

    # 10. Disposições gerais
    h1(doc, "10. Disposições gerais")
    para(
        doc,
        "10.1. Este Termo constitui o acordo integral entre as Partes quanto ao "
        "seu objeto, prevalecendo sobre tratativas anteriores. Alterações somente "
        "por escrito.",
    )
    para(
        doc,
        "10.2. A eventual tolerância quanto ao descumprimento de qualquer cláusula "
        "não implica renúncia de direitos.",
    )
    para(
        doc,
        "10.3. A nulidade parcial de qualquer disposição não afeta as demais.",
    )
    para(
        doc,
        "10.4. Fica eleito o foro da Comarca de Belo Horizonte/MG, com renúncia a "
        "qualquer outro, por mais privilegiado que seja, para dirimir controvérsias "
        "oriundas deste Termo, aplicando-se a legislação brasileira.",
    )
    para(
        doc,
        "10.5. As Partes reconhecem ter lido e compreendido todas as cláusulas, "
        "assinando o presente em duas vias de igual teor, na data indicada abaixo.",
    )

    # Assinaturas
    h1(doc, "11. Assinaturas")
    para(doc, f"Local e data: ________________________, ____/____/________",
         space_after=16)

    # Two signature blocks
    table = doc.add_table(rows=1, cols=2)
    table.autofit = True
    for i, (title, lines) in enumerate([
        (
            "Pela COMODANTE — MIRA",
            "________________________________\n"
            "MIRA — Mapeamento e Robótica Ltda.\n"
            "Nome: Gilmar Pereira da Cruz Júnior\n"
            "Cargo: Representante legal\n"
            "CPF: ___________________________",
        ),
        (
            "Pelo COMODATÁRIO",
            "________________________________\n"
            "Nome: ___________________________\n"
            "CPF/CNPJ: _______________________\n"
            "Cargo / função: _________________\n"
            "E-mail: _________________________",
        ),
    ]):
        cell = table.cell(0, i)
        set_cell_margins(cell, 80, 80, 60, 60)
        p = cell.paragraphs[0]
        r = p.add_run(title)
        set_run_font(r, size=9, bold=True, color=DARK)
        for line in lines.split("\n"):
            p2 = cell.add_paragraph()
            p2.paragraph_format.space_after = Pt(2)
            r2 = p2.add_run(line)
            set_run_font(r2, size=9, color=DARK)

    para(doc, "", space_after=12)
    para(
        doc,
        "Testemunhas (opcional):",
        bold=True, space_after=8,
    )
    field_row(doc, "1. Nome / CPF:")
    field_row(doc, "2. Nome / CPF:")

    para(
        doc,
        "Contato MIRA: contato@mirarobotica.com · www.mirarobotica.com",
        size=9, align="center", space_before=18, color=GRAY,
    )

    doc.save(str(OUT))
    print(f"OK {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    build()
